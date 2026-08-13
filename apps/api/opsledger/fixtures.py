from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal

import pandas as pd
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .schemas import CaseCreate


@dataclass(frozen=True)
class FixtureDocument:
    document_type: str
    filename: str
    mime_type: str
    content: bytes


@dataclass(frozen=True)
class CaseFixture:
    reference: str
    profile: CaseCreate
    documents: tuple[FixtureDocument, ...]


def _evidence_pdf(title: str, reference: str, rows: list[tuple[str, str]]) -> bytes:
    output = io.BytesIO()
    page = canvas.Canvas(output, pagesize=A4, invariant=1)
    width, height = A4
    page.setTitle(f"{reference} {title}")
    page.setFillColor(colors.HexColor("#173f37"))
    page.rect(0, height - 92, width, 92, fill=True, stroke=False)
    page.setFillColor(colors.white)
    page.setFont("Helvetica-Bold", 10)
    page.drawString(48, height - 42, "OPSLEDGER AI · SYNTHETIC EVIDENCE")
    page.setFont("Helvetica-Bold", 20)
    page.drawString(48, height - 70, title)
    page.setFillColor(colors.HexColor("#1d2925"))
    page.setFont("Helvetica", 9)
    page.drawString(48, height - 122, f"Case reference: {reference}")
    page.setStrokeColor(colors.HexColor("#d4ddd8"))
    page.line(48, height - 136, width - 48, height - 136)
    y = height - 172
    for label, value in rows:
        page.setFillColor(colors.HexColor("#66736e"))
        page.setFont("Helvetica-Bold", 8)
        page.drawString(48, y, label.upper())
        page.setFillColor(colors.HexColor("#1d2925"))
        page.setFont("Helvetica", 11)
        page.drawString(48, y - 18, f"{label}: {value}")
        y -= 60
    page.setFillColor(colors.HexColor("#66736e"))
    page.setFont("Helvetica", 8)
    page.drawString(48, 40, "Synthetic buildathon fixture. Not an official record.")
    page.save()
    return output.getvalue()


def _registration_pdf(
    reference: str,
    legal_name: str,
    registration_number: str,
    jurisdiction: str,
) -> bytes:
    return _evidence_pdf(
        "Registration evidence",
        reference,
        [
            ("Legal Business Name", legal_name),
            ("Registration Number", registration_number),
            ("Jurisdiction", jurisdiction),
            ("Registered On", "2022-04-18"),
        ],
    )


def _ownership_pdf(reference: str, legal_name: str, owner: str) -> bytes:
    return _evidence_pdf(
        "Ownership declaration",
        reference,
        [
            ("Legal Business Name", legal_name),
            ("Ownership Declared", "Yes"),
            ("Owner Name", owner),
            ("Signed On", "2026-07-29"),
        ],
    )


def _revenue_workbook(
    reference: str,
    amounts: list[Decimal],
    as_of_date: date,
    currency: str,
) -> bytes:
    periods = pd.date_range(end=as_of_date, periods=len(amounts), freq="ME")
    frame = pd.DataFrame(
        {
            "period": [period.strftime("%Y-%m") for period in periods],
            "revenue": [float(amount) for amount in amounts],
            "currency": [currency] * len(amounts),
            "as_of_date": [as_of_date.isoformat()] * len(amounts),
            "declared_total": [float(sum(amounts))] + [None] * (len(amounts) - 1),
            "case_reference": [reference] * len(amounts),
        }
    )
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Revenue")
        sheet = writer.book["Revenue"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="173F37")
        for column, width in {"A": 14, "B": 16, "C": 12, "D": 16, "E": 18, "F": 20}.items():
            sheet.column_dimensions[column].width = width
    return output.getvalue()


def _bank_csv(
    reference: str,
    as_of_date: date,
    currency: str,
    opening_balance: Decimal,
) -> bytes:
    entries = [
        ("Customer receipts", Decimal("18450.00")),
        ("Supplier payment", Decimal("-7280.00")),
        ("Payroll", Decimal("-9400.00")),
        ("Card settlements", Decimal("12620.00")),
    ]
    balance = opening_balance
    rows: list[dict[str, object]] = []
    first_date = as_of_date - timedelta(days=len(entries) - 1)
    for offset, (description, amount) in enumerate(entries):
        balance += amount
        rows.append(
            {
                "date": (first_date + timedelta(days=offset)).isoformat(),
                "description": f"{description} · {reference}",
                "amount": str(amount),
                "balance": str(balance),
                "currency": currency,
            }
        )
    return pd.DataFrame(rows).to_csv(index=False).encode("utf-8")


def _documents(
    *,
    reference: str,
    intake_name: str,
    registration_name: str,
    registration_number: str,
    jurisdiction: str,
    owner: str,
    revenue: list[Decimal],
    financial_as_of: date,
    currency: str,
    include_ownership: bool = True,
) -> tuple[FixtureDocument, ...]:
    documents = [
        FixtureDocument(
            "REGISTRATION_EVIDENCE",
            f"{reference.lower()}-registration.pdf",
            "application/pdf",
            _registration_pdf(
                reference,
                registration_name,
                registration_number,
                jurisdiction,
            ),
        ),
        FixtureDocument(
            "REVENUE_STATEMENT",
            f"{reference.lower()}-revenue.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _revenue_workbook(reference, revenue, financial_as_of, currency),
        ),
        FixtureDocument(
            "BANK_STATEMENT",
            f"{reference.lower()}-bank.csv",
            "text/csv",
            _bank_csv(reference, financial_as_of, currency, Decimal("94000.00")),
        ),
    ]
    if include_ownership:
        documents.append(
            FixtureDocument(
                "OWNERSHIP_DECLARATION",
                f"{reference.lower()}-ownership.pdf",
                "application/pdf",
                _ownership_pdf(reference, intake_name, owner),
            )
        )
    return tuple(documents)


def demo_fixtures() -> tuple[CaseFixture, ...]:
    case_a_name = "Island Harvest Foods Ltd."
    case_b_name = "Blue Shore Repairs"
    case_c_name = "Caribbean Green Logistics Ltd."
    return (
        CaseFixture(
            reference="OPS-2026-0001",
            profile=CaseCreate(
                legal_business_name=case_a_name,
                trading_name="Island Harvest",
                registration_number="ABR-4421-A",
                jurisdiction="Antigua and Barbuda",
                industry="Food manufacturing",
                requested_amount=Decimal("85000.00"),
                currency="XCD",
                funding_purpose="Purchase a cold-room unit and increase regional distribution.",
                annual_revenue=Decimal("480000.00"),
                contact_name="Alana Joseph",
                contact_email="alana@islandharvest.example",
            ),
            documents=_documents(
                reference="OPS-2026-0001",
                intake_name=case_a_name,
                registration_name=case_a_name,
                registration_number="ABR-4421-A",
                jurisdiction="Antigua and Barbuda",
                owner="Alana Joseph",
                revenue=[Decimal("80000.00")] * 6,
                financial_as_of=date(2026, 7, 31),
                currency="XCD",
            ),
        ),
        CaseFixture(
            reference="OPS-2026-0002",
            profile=CaseCreate(
                legal_business_name=case_b_name,
                trading_name=None,
                registration_number="LCR-7319-R",
                jurisdiction="Saint Lucia",
                industry="Marine equipment repair",
                requested_amount=Decimal("42000.00"),
                currency="XCD",
                funding_purpose="Replace diagnostic tools and add one service vehicle.",
                annual_revenue=Decimal("210000.00"),
                contact_name="Micah James",
                contact_email="micah@blueshore.example",
            ),
            documents=_documents(
                reference="OPS-2026-0002",
                intake_name=case_b_name,
                registration_name=case_b_name,
                registration_number="LCR-7319-R",
                jurisdiction="Saint Lucia",
                owner="Micah James",
                revenue=[Decimal("35000.00")] * 6,
                financial_as_of=date(2025, 10, 31),
                currency="XCD",
                include_ownership=False,
            ),
        ),
        CaseFixture(
            reference="OPS-2026-0003",
            profile=CaseCreate(
                legal_business_name=case_c_name,
                trading_name="GreenRoute Caribbean",
                registration_number="ABR-4421-A",
                jurisdiction="Antigua and Barbuda",
                industry="Low-emission freight and logistics",
                requested_amount=Decimal("125000.00"),
                currency="XCD",
                funding_purpose="Acquire two electric delivery vans for inter-island contracts.",
                annual_revenue=Decimal("720000.00"),
                contact_name="Zara Clarke",
                contact_email="zara@greenroute.example",
            ),
            documents=_documents(
                reference="OPS-2026-0003",
                intake_name=case_c_name,
                registration_name="Caribbean Green Freight Services Ltd.",
                registration_number="ABR-4421-A",
                jurisdiction="Antigua and Barbuda",
                owner="Zara Clarke",
                revenue=[Decimal("120000.00")] * 6,
                financial_as_of=date(2026, 7, 31),
                currency="XCD",
            ),
        ),
    )
